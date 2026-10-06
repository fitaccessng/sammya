from io import BytesIO
import os

from flask import url_for
from openpyxl import Workbook

from app.factory import create_app
from app.models import (
    BOQImport,
    BOQItem,
    Project,
    ProjectDocument,
    ProjectStaff,
    User,
    db,
)


def _create_app_and_db():
    app = create_app('development')
    app.config.update(
        TESTING=True,
        WTF_CSRF_ENABLED=False,
        SQLALCHEMY_DATABASE_URI='sqlite:///:memory:',
    )
    with app.app_context():
        db.drop_all()
        db.create_all()
    return app


def test_qs_manager_can_upload_to_assigned_project():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Alpha', budget=1000)
        db.session.add(project)
        db.session.flush()

        qs_user = User(name='QS Manager', email='qs@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add(qs_user)
        db.session.flush()

        db.session.add(ProjectStaff(user_id=qs_user.id, project_id=project.id, role='QS Manager', is_active=True))
        db.session.commit()
        qs_user_id = qs_user.id
        project_id = project.id

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(qs_user_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/documents/upload',
            data={
                'title': 'BOQ Upload',
                'document_type': 'XLSX',
                'file': (BytesIO(b'not really an xlsx file'), 'LABOR ONLY FOR GATE HOUSE.xlsx'),
            },
            content_type='multipart/form-data',
        )

        assert response.status_code == 302
        assert f'/projects/{project_id}/documents' in response.headers['Location']


def test_qs_boq_upload_preserves_custom_file_columns(tmp_path):
    app = _create_app_and_db()
    app.config['PROJECT_UPLOAD_FOLDER'] = str(tmp_path)
    with app.app_context():
        project = Project(name='Custom BOQ Project', budget=1000)
        qs_user = User(name='QS Manager', email='custom-boq@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add_all([project, qs_user])
        db.session.flush()
        db.session.add(ProjectStaff(
            user_id=qs_user.id,
            project_id=project.id,
            role='QS Manager',
            is_active=True,
        ))
        db.session.commit()
        project_id = project.id
        user_id = qs_user.id
        with app.test_request_context():
            upload_url = url_for('qs_boq.upload_boq', project_id=project_id)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True

        response = client.post(
            upload_url,
            data={
                'file': (BytesIO(
                    b'Zone,Work Package,Measured Area,Rate NGN,Checked By\n'
                    b'North,Concrete slab,24.5,18500,QS Team\n'
                ), 'custom-layout.csv'),
            },
            content_type='multipart/form-data',
        )
        page_response = client.get(f'/qs/project/{project_id}/boq')

    assert response.status_code == 200
    assert response.get_json()['created'] == 1
    assert page_response.status_code == 200
    assert b'Imported BOQ Rows' in page_response.data
    assert b'Measured Area' in page_response.data
    assert b'24.5' in page_response.data
    with app.app_context():
        boq_import = BOQImport.query.one()
        imported_sheet = boq_import.sheets[0]
        imported_row = imported_sheet.rows[0]
        assert imported_sheet.headers == [
            'Zone', 'Work Package', 'Measured Area', 'Rate NGN', 'Checked By',
        ]
        assert imported_row.values == ['North', 'Concrete slab', 24.5, 18500, 'QS Team']
        assert BOQItem.query.count() == 0


def test_qs_project_upload_link_and_document_survives_missing_local_file(tmp_path):
    app = _create_app_and_db()
    app.config['PROJECT_UPLOAD_FOLDER'] = str(tmp_path)
    with app.app_context():
        project = Project(name='Persistent BOQ Project', budget=1000)
        qs_user = User(name='QS Manager', email='persistent-boq@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add_all([project, qs_user])
        db.session.flush()
        db.session.add(ProjectStaff(
            user_id=qs_user.id,
            project_id=project.id,
            role='QS Manager',
            is_active=True,
        ))
        db.session.commit()
        project_id = project.id
        user_id = qs_user.id
        with app.test_request_context():
            upload_url = url_for('qs_boq.upload_boq', project_id=project_id)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True

        project_page = client.get(f'/qs/project/{project_id}')
        assert project_page.status_code == 200
        assert b'Upload BOQ' in project_page.data

        response = client.post(
            upload_url,
            data={'file': (BytesIO(b'Item,Description\n1,Test row\n'), 'boq.csv')},
            content_type='multipart/form-data',
        )
        assert response.status_code == 200

        with app.app_context():
            document = ProjectDocument.query.filter_by(project_id=project_id).one()
            document_id = document.id
            local_path = document.file_path
            assert document.file_data == b'Item,Description\n1,Test row\n'

        os.remove(local_path)
        viewed = client.get(f'/projects/documents/{document_id}/view')

    assert viewed.status_code == 200
    assert viewed.data == b'Item,Description\n1,Test row\n'


def test_qs_boq_upload_preserves_all_excel_sheets(tmp_path):
    app = _create_app_and_db()
    app.config['PROJECT_UPLOAD_FOLDER'] = str(tmp_path)
    with app.app_context():
        project = Project(name='Workbook BOQ Project', budget=1000)
        qs_user = User(name='QS Manager', email='workbook-boq@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add_all([project, qs_user])
        db.session.flush()
        db.session.add(ProjectStaff(
            user_id=qs_user.id,
            project_id=project.id,
            role='QS Manager',
            is_active=True,
        ))
        db.session.commit()
        project_id = project.id
        user_id = qs_user.id
        with app.test_request_context():
            upload_url = url_for('qs_boq.upload_boq', project_id=project_id)

    workbook = Workbook()
    workbook.active.title = 'Measured Works'
    workbook.active.append(['Location', 'Trade Description', 'Measured Area'])
    workbook.active.append(['Block A', 'Floor screed', 75.5])
    notes_sheet = workbook.create_sheet('Commercial Notes')
    notes_sheet.append(['Clause Reference', 'Note'])
    notes_sheet.append(['C-14', 'Rate subject to review'])
    workbook_file = BytesIO()
    workbook.save(workbook_file)
    workbook_file.seek(0)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True
        response = client.post(
            upload_url,
            data={'file': (workbook_file, 'multi-sheet-boq.xlsx')},
            content_type='multipart/form-data',
        )

    assert response.status_code == 200
    assert response.get_json()['created'] == 2
    assert response.get_json()['sheets'] == 2
    with app.app_context():
        boq_import = BOQImport.query.one()
        sheets = {sheet.sheet_name: sheet for sheet in boq_import.sheets}
        assert set(sheets) == {'Measured Works', 'Commercial Notes'}
        assert sheets['Measured Works'].headers == ['Location', 'Trade Description', 'Measured Area']
        assert sheets['Measured Works'].rows[0].values == ['Block A', 'Floor screed', 75.5]
        assert sheets['Commercial Notes'].rows[0].values == ['C-14', 'Rate subject to review']


def test_material_schedule_upload_uses_supported_boq_fields_and_renders(tmp_path):
    app = _create_app_and_db()
    app.config['PROJECT_UPLOAD_FOLDER'] = str(tmp_path)
    with app.app_context():
        project = Project(name='Material Schedule Project', budget=1000)
        qs_user = User(name='QS Manager', email='schedule-boq@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add_all([project, qs_user])
        db.session.flush()
        db.session.add(ProjectStaff(
            user_id=qs_user.id,
            project_id=project.id,
            role='QS Manager',
            is_active=True,
        ))
        db.session.commit()
        project_id = project.id
        user_id = qs_user.id
        with app.test_request_context():
            upload_url = url_for('qs_boq.upload_material_schedule', project_id=project_id)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True
        response = client.post(
            upload_url,
            data={
                'file': (BytesIO(
                    b'Item ID,Description,Material Qty,Material Unit,Material Rate,Material Total,Labour Qty,Labour Unit,Labour Rate,Labour Total,Grand Total\n'
                    b'P-01,Painting,5,L,20,100,2,HRS,15,30,130\n'
                ), 'material-schedule.csv'),
            },
            content_type='multipart/form-data',
        )
        material_response = client.get(f'/qs/project/{project_id}/material-schedule')
        project_response = client.get(f'/qs/project/{project_id}')

    assert response.status_code == 200
    assert response.get_json()['created'] == 2
    assert material_response.status_code == 200
    assert b'Materials' in material_response.data
    assert b'Labour' in material_response.data
    assert project_response.status_code == 200
    with app.app_context():
        items = BOQItem.query.filter_by(project_id=project_id).order_by(BOQItem.item_no).all()
        assert [item.item_no for item in items] == ['P-01-L', 'P-01-M']
        assert all(item.created_by == user_id for item in items)


def test_qs_manager_cannot_upload_to_unassigned_project():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Beta', budget=1000)
        db.session.add(project)
        db.session.flush()

        qs_user = User(name='QS Manager', email='qs2@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add(qs_user)
        db.session.commit()
        qs_user_id = qs_user.id
        project_id = project.id

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(qs_user_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/documents/upload',
            data={
                'file': (BytesIO(b'bad'), 'Other.xlsx'),
            },
            content_type='multipart/form-data',
        )

        assert response.status_code == 403


def test_admin_can_upload_to_project():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Gamma', budget=1000)
        db.session.add(project)
        db.session.flush()

        admin_user = User(name='Admin User', email='admin@example.com', role='admin')
        admin_user.set_password('password')
        db.session.add(admin_user)
        db.session.commit()
        admin_user_id = admin_user.id
        project_id = project.id

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(admin_user_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/documents/upload',
            data={
                'title': 'Admin Upload',
                'document_type': 'PDF',
                'file': (BytesIO(b'pdf contents'), 'sample.pdf'),
            },
            content_type='multipart/form-data',
        )

        assert response.status_code == 302
        assert f'/projects/{project_id}/documents' in response.headers['Location']


def test_authenticated_project_user_can_upload_unrestricted_extension():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Delta', budget=2000)
        db.session.add(project)
        db.session.flush()

        engineer = User(name='Engineer', email='engineer@example.com', role='engineer')
        engineer.set_password('password')
        db.session.add(engineer)
        db.session.flush()

        db.session.add(
            ProjectStaff(user_id=engineer.id, project_id=project.id, role='Engineer', is_active=True)
        )
        db.session.commit()
        project_id = project.id
        engineer_id = engineer.id

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(engineer_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/documents/upload',
            data={
                'title': 'DWG Upload',
                'document_type': 'DWG',
                'file': (BytesIO(b'fake dwg content'), 'site_layout.dwg'),
            },
            content_type='multipart/form-data',
        )

        assert response.status_code == 302
        assert f'/projects/{project_id}/documents' in response.headers['Location']


def test_unassigned_user_is_denied_for_any_extension():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Epsilon', budget=3000)
        db.session.add(project)
        db.session.flush()
        project_id = project.id

        site_manager = User(name='Site Manager', email='site@example.com', role='site_manager')
        site_manager.set_password('password')
        db.session.add(site_manager)
        db.session.commit()
        site_manager_id = site_manager.id

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(site_manager_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/documents/upload',
            data={
                'title': 'Blocked Upload',
                'file': (BytesIO(b'bad'), 'blocked.zip'),
            },
            content_type='multipart/form-data',
        )

        assert response.status_code == 403


def test_project_manager_imports_boq_with_external_headers_into_project_model():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='BOQ Import Project', budget=500000)
        manager = User(name='Project Manager', email='boq-manager@example.com', role='project_manager')
        manager.set_password('password')
        db.session.add_all([project, manager])
        db.session.commit()
        project_id = project.id
        manager_id = manager.id

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Item No', 'Work Description', 'UOM', 'Measured Qty', 'Unit Rate', 'Value'])
    sheet.append(['A-01', 'Concrete works', 'm3', 10, 5000, 50000])
    file_content = BytesIO()
    workbook.save(file_content)
    file_content.seek(0)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(manager_id)
            session['_fresh'] = True

        response = client.post(
            f'/projects/{project_id}/boq/import',
            data={'boq_file': (file_content, 'contractor-boq.xlsx')},
            content_type='multipart/form-data',
        )

        assert response.status_code == 302
        assert f'/projects/{project_id}' in response.headers['Location']

    with app.app_context():
        imported = BOQItem.query.filter_by(project_id=project_id).one()
        assert imported.description == 'Concrete works'
        assert imported.unit == 'm3'
        assert float(imported.quantity) == 10
        assert float(imported.unit_rate) == 5000
        assert float(imported.amount) == 50000


def test_qs_boq_page_and_add_work_without_archive_tables():
    app = _create_app_and_db()
    with app.app_context():
        project = Project(name='Legacy BOQ Project', budget=1000)
        qs_user = User(name='QS Manager', email='legacy-boq@example.com', role='qs_manager')
        qs_user.set_password('password')
        db.session.add_all([project, qs_user])
        db.session.flush()
        db.session.add(ProjectStaff(
            user_id=qs_user.id,
            project_id=project.id,
            role='QS Manager',
            is_active=True,
        ))
        db.session.commit()
        project_id = project.id
        user_id = qs_user.id

        from app.models import BOQImportRow, BOQImportSheet
        BOQImportRow.__table__.drop(bind=db.engine, checkfirst=True)
        BOQImportSheet.__table__.drop(bind=db.engine, checkfirst=True)
        BOQImport.__table__.drop(bind=db.engine, checkfirst=True)

    with app.test_client() as client:
        with client.session_transaction() as session:
            session['_user_id'] = str(user_id)
            session['_fresh'] = True

        page = client.get(f'/qs/project/{project_id}/boq')
        assert page.status_code == 200
        assert b'Uploaded-file archive storage is not initialized yet' in page.data

        response = client.post(
            f'/qs/project/{project_id}/boq/add',
            json={'description': 'Concrete slab', 'quantity': 10, 'unit': 'm3', 'unit_rate': 250},
        )
        assert response.status_code == 200
        assert response.get_json()['success'] is True

    with app.app_context():
        item = BOQItem.query.filter_by(project_id=project_id).one()
        assert item.description == 'Concrete slab'
        assert item.created_by == user_id

        from app.models import BOQImportRow, BOQImportSheet, ensure_boq_import_tables
        from sqlalchemy import inspect
        ensure_boq_import_tables()
        ensure_boq_import_tables()
        inspector = inspect(db.engine)
        assert inspector.has_table('boq_import')
        assert inspector.has_table('boq_import_sheet')
        assert inspector.has_table('boq_import_row')
